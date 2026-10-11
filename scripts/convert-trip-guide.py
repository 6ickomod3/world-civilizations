#!/usr/bin/env python3
"""Read the Europe inventory without modifying it; emit public guide JSON.

Run with the bundled Python reported by load_workspace_dependencies:
    python3 scripts/convert-trip-guide.py
    python3 scripts/convert-trip-guide.py --check

Each positive field map below documents a source-header -> public-field mapping.
Editable Decision and Personal notes columns are intentionally never copied.
Free-text historical periods stay free text: this converter invents no dates,
coordinates, confirmed visits, route commitments, or new candidate identities.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from collections import Counter
from pathlib import Path
from urllib.parse import urlsplit

import openpyxl


REPO = Path(__file__).resolve().parents[1]
HEADER_ROW = 5
EXPECTED_COUNTS = {"Overview": 15, "Places": 354, "Connections": 112, "Coverage gaps": 13}
HEADERS = {
    "Overview": (
        "Leg", "Journey", "Countries / region", "Guiding question",
        "Civilization storyline", "Geographic backbone", "10-day feasibility",
        "Category coverage", "Connected legs", "Candidates",
    ),
    "Places": (
        "ID", "Leg", "Decision", "Place / experience", "City or area", "Country",
        "Cluster", "Category", "Era / period", "Candidate status", "What to see or do",
        "Why it matters", "Connects to", "Geography and route notes", "Source URLs",
        "Personal notes",
    ),
    "Connections": (
        "ID", "Theme", "From leg(s)", "To leg(s)", "Connection or coverage gap",
        "What to compare", "Places / regions", "Notes",
    ),
    "Coverage gaps": (
        "ID", "Theme", "Region", "Why it matters", "Places and collections to consider",
        "Relation to the 15 legs", "Geography and route notes", "Source URLs",
        "Decision", "Personal notes",
    ),
}
FIELD_MAPS = {
    "Overview": {
        "Leg": "leg", "Journey": "title", "Countries / region": "countriesRegion",
        "Guiding question": "guidingQuestion", "Civilization storyline": "civilizationStoryline",
        "Geographic backbone": "geographicBackbone", "10-day feasibility": "tenDayFeasibility",
        "Category coverage": "categoryCoverage", "Connected legs": "connectedLegsText",
        "Candidates": "candidateCount",
    },
    "Places": {
        "ID": "id", "Leg": "leg", "Place / experience": "name", "City or area": "cityArea",
        "Country": "country", "Cluster": "cluster", "Category": "category",
        "Era / period": "eraPeriod", "Candidate status": "candidateStatus",
        "What to see or do": "whatToSee", "Why it matters": "whyItMatters",
        "Connects to": "connectsTo", "Geography and route notes": "routeNotes",
        "Source URLs": "sourceUrlsText",
    },
    "Connections": {
        "ID": "id", "Theme": "theme", "From leg(s)": "fromLegsText",
        "To leg(s)": "toLegsText", "Connection or coverage gap": "description",
        "What to compare": "whatToCompare", "Places / regions": "placesRegions", "Notes": "notes",
    },
    "Coverage gaps": {
        "ID": "id", "Theme": "theme", "Region": "region", "Why it matters": "whyItMatters",
        "Places and collections to consider": "placesToConsider",
        "Relation to the 15 legs": "relationToLegs",
        "Geography and route notes": "routeNotes", "Source URLs": "sourceUrlsText",
    },
}
PRIVATE_FIELDS = {"Decision", "Personal notes"}
COUNTRY_CODES = {
    "Greece": ["GRC"], "Italy": ["ITA"], "Vatican City": ["VAT"], "Türkiye": ["TUR"],
    "Denmark": ["DNK"], "Sweden": ["SWE"], "Norway": ["NOR"], "Spain": ["ESP"],
    "France": ["FRA"], "Portugal": ["PRT"], "Belgium": ["BEL"], "Netherlands": ["NLD"],
    "Czechia": ["CZE"], "Austria": ["AUT"], "Hungary": ["HUN"], "Switzerland": ["CHE"],
    "Switzerland–Italy": ["CHE", "ITA"], "United Kingdom": ["GBR"], "Germany": ["DEU"],
    "Poland": ["POL"], "Slovenia": ["SVN"], "Croatia": ["HRV"],
    "Bosnia and Herzegovina": ["BIH"],
}
PERSON_NAME = r"[A-Z][\w'’.-]*(?:\s+[A-Z][\w'’.-]*)*"
ANONYMOUS_REPLACEMENTS = (
    (
        re.compile(rf"\b{PERSON_NAME}\s+has\s+already\s+visited[^.]*\."),
        "Prior visits should be reviewed before choosing stops. This inventory does not assume which individual sites have been visited.",
    ),
    (
        re.compile(rf"\b{PERSON_NAME}\s+has\s+visited[^.]*\."),
        "Prior visits should be reviewed before choosing stops; this particular site is not assumed visited.",
    ),
    (
        re.compile(r"\bDo not assume a repeat ([A-Za-z][A-Za-z -]*?) trip is needed until [^.]*actual prior itinerary is mapped\."),
        r"Prior visits should be reviewed before deciding whether another \1 journey is needed.",
    ),
)
PERSONAL_COMMENTARY = re.compile(
    rf"\b{PERSON_NAME}\s+(?:has\s+(?:already\s+)?visited|(?:wants?|plans?)\s+to\s+visit)\b"
    rf"|\b{PERSON_NAME}['’]s\s+(?:actual\s+)?prior\s+(?:visits?|itinerary)\b"
    r"|\b(?:the\s+)?user['’]s\s+(?:actual\s+)?(?:prior\s+)?(?:visits?|itinerary)\b"
)
LEG_REFERENCES = re.compile(r"\blegs?\s*(\d+(?:(?:\s*[,/&–-]\s*|\s*and\s*)\d+)*)", re.I)


def public_value(value):
    if not isinstance(value, str):
        return value
    for original, anonymous in ANONYMOUS_REPLACEMENTS:
        value = original.sub(anonymous, value)
    if PERSONAL_COMMENTARY.search(value):
        raise ValueError("Unrecognized named personal commentary; add an explicit neutral replacement.")
    return value


def read_tables(path):
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        if set(workbook.sheetnames) != set(HEADERS):
            raise ValueError(f"Unexpected sheets: {workbook.sheetnames}")
        tables = {}
        for name, expected in HEADERS.items():
            rows = iter(workbook[name].iter_rows(values_only=True))
            for _ in range(HEADER_ROW - 1):
                next(rows)
            actual = tuple(next(rows))
            if actual != expected:
                raise ValueError(f"{name}: unexpected headers {actual!r}")
            records = []
            for row_number, values in enumerate(rows, HEADER_ROW + 1):
                if not any(value is not None for value in values):
                    continue
                if len(values) != len(expected):
                    raise ValueError(f"{name} row {row_number}: unexpected column count")
                records.append((row_number, dict(zip(expected, values))))
            if len(records) != EXPECTED_COUNTS[name]:
                raise ValueError(f"{name}: found {len(records)} records, expected {EXPECTED_COUNTS[name]}")
            tables[name] = records
        return tables
    finally:
        workbook.close()


def mapped_record(sheet, row_number, row):
    return {
        **{field: public_value(row[header]) for header, field in FIELD_MAPS[sheet].items()},
        "sourceRow": row_number,
    }


def strict_leg_list(text):
    if not isinstance(text, str) or not re.fullmatch(r"\s*\d+(?:\s*,\s*\d+)*\s*", text):
        raise ValueError(f"Unexpected leg list: {text!r}")
    legs = [int(value.strip()) for value in text.split(",")]
    if any(leg not in range(1, 16) for leg in legs):
        raise ValueError(f"Unknown journey leg in {text!r}")
    return legs


def related_legs(text):
    legs = []
    for match in LEG_REFERENCES.finditer(text or ""):
        sequence = match.group(1)
        for token in re.split(r"\s*(?:,|/|&|and)\s*", sequence, flags=re.I):
            if re.search(r"[–-]", token):
                start, end = (int(value.strip()) for value in re.split(r"[–-]", token))
                values = range(start, end + 1)
            else:
                values = [int(token.strip())]
            for leg in values:
                if leg not in range(1, 16):
                    raise ValueError(f"Unknown leg reference in {text!r}")
                if leg not in legs:
                    legs.append(leg)
    return legs


def source_urls(text):
    urls = [line.strip() for line in text.splitlines() if line.strip()]
    if not urls:
        raise ValueError("A source URL field has no URLs")
    for url in urls:
        parsed = urlsplit(url)
        if parsed.scheme not in {"http", "https"} or not parsed.netloc or re.search(r"\s", url):
            raise ValueError(f"Malformed source URL: {url!r}")
    return urls


def country_codes(country):
    if country not in COUNTRY_CODES:
        raise ValueError(f"Unmapped country reference: {country!r}")
    return list(COUNTRY_CODES[country])


def require_unique(records, key, label):
    values = [record[key] for record in records]
    if any(value is None for value in values) or len(set(values)) != len(values):
        raise ValueError(f"{label}: missing or duplicate {key}")


def convert(tables):
    journeys, places, connections, gaps = [], [], [], []
    for row_number, row in tables["Overview"]:
        item = mapped_record("Overview", row_number, row)
        item["id"] = f"journey-{item['leg']:02d}"
        item["categories"] = [part.strip() for part in item["categoryCoverage"].split(",")]
        item["connectedLegs"] = strict_leg_list(item["connectedLegsText"])
        codes = []
        for country in item["countriesRegion"].split(","):
            for code in country_codes(country.strip()):
                if code not in codes:
                    codes.append(code)
        item["countryCodes"] = codes
        journeys.append(item)
    for row_number, row in tables["Places"]:
        item = mapped_record("Places", row_number, row)
        item["relatedLegs"] = related_legs(item["connectsTo"])
        item["sourceUrls"] = source_urls(item["sourceUrlsText"])
        item["countryCodes"] = country_codes(item["country"])
        identifier = re.fullmatch(r"L?(\d{2})-\d{2}", item["id"])
        if item["leg"] not in range(1, 16) or not identifier or int(identifier.group(1)) != item["leg"]:
            raise ValueError(f"Inconsistent place ID / leg: {item['id']!r}")
        places.append(item)
    for row_number, row in tables["Connections"]:
        item = mapped_record("Connections", row_number, row)
        item["fromLegs"] = strict_leg_list(item["fromLegsText"])
        item["toLegs"] = strict_leg_list(item["toLegsText"])
        connections.append(item)
    for row_number, row in tables["Coverage gaps"]:
        item = mapped_record("Coverage gaps", row_number, row)
        item["relatedLegs"] = related_legs(item["relationToLegs"])
        item["sourceUrls"] = source_urls(item["sourceUrlsText"])
        gaps.append(item)

    for records, label in [(journeys, "journeys"), (places, "places"), (connections, "connections"), (gaps, "coverage gaps")]:
        require_unique(records, "id", label)
    require_unique(journeys, "leg", "journeys")
    if [item["leg"] for item in journeys] != list(range(1, 16)):
        raise ValueError("Overview journey sequence does not cover legs 1–15")
    place_counts = Counter(item["leg"] for item in places)
    for journey in journeys:
        if journey["candidateCount"] != place_counts[journey["leg"]]:
            raise ValueError(f"Leg {journey['leg']}: Overview candidate count disagrees with Places")
    for sheet, records in [("Places", places), ("Connections", connections), ("Coverage gaps", gaps)]:
        if [row["ID"] for _, row in tables[sheet]] != [item["id"] for item in records]:
            raise ValueError(f"{sheet}: original ID order was not preserved")
        for (_, original), converted in zip(tables[sheet], records):
            for header, field in FIELD_MAPS[sheet].items():
                if converted[field] != public_value(original[header]):
                    raise ValueError(f"{sheet}: source field changed: {header}")
    for mapping in FIELD_MAPS.values():
        if PRIVATE_FIELDS.intersection(mapping):
            raise ValueError("Private editable field entered the public allowlist")

    metadata = {
        "schemaVersion": 1,
        "sourceWorkbook": "europe_civilizations_inventory.xlsx",
        "sourceHeaderRow": HEADER_ROW,
        "counts": {"journeys": len(journeys), "places": len(places), "connections": len(connections), "coverageGaps": len(gaps)},
        "headerMap": FIELD_MAPS,
        "derivedFields": {
            "id": "Journey IDs only are generated from leg number; place, connection and gap IDs are original.",
            "sourceRow": "Original one-based workbook row.",
            "countryCodes": "ISO3 references for present-day geography, including both countries for Switzerland–Italy.",
            "categories": "Comma-separated Overview category coverage, with original text also preserved.",
            "connectedLegs": "Numeric Overview leg references, with original text also preserved.",
            "fromLegs": "Numeric connection origins, with original text also preserved.",
            "toLegs": "Numeric connection destinations, with original text also preserved.",
            "relatedLegs": "Leg references parsed from narrative links, with original narrative also preserved.",
            "sourceUrls": "Nonblank source URL lines, with original cell text also preserved.",
        },
        "anonymizedCommentaryCells": sum(
            public_value(row[header]) != row[header]
            for sheet, records in tables.items()
            for _, row in records
            for header in FIELD_MAPS[sheet]
        ),
        "privateEditableFieldsExcluded": True,
        "confirmedVisitsInferred": False,
    }
    return {"journeys": journeys, "connections": connections, "coverageGaps": gaps, "metadata": metadata}, places


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=REPO / "europe/europe_civilizations_inventory.xlsx")
    parser.add_argument("--output-dir", type=Path, default=REPO / "europe")
    parser.add_argument("--check", action="store_true", help="Compare regenerated content with existing JSON; write nothing.")
    args = parser.parse_args()
    before = hashlib.sha256(args.input.read_bytes()).hexdigest()
    tables = read_tables(args.input)
    journeys, places = convert(tables)
    outputs = {"journeys-source.json": journeys, "places-source.json": places}
    rendered = {name: json.dumps(payload, ensure_ascii=False, indent=2, allow_nan=False) + "\n" for name, payload in outputs.items()}
    after = hashlib.sha256(args.input.read_bytes()).hexdigest()
    if before != after:
        raise ValueError("Source workbook changed during conversion")
    if args.check:
        for name, content in rendered.items():
            if (args.output_dir / name).read_text(encoding="utf-8") != content:
                raise ValueError(f"Output differs from reproducible source conversion: {name}")
    else:
        args.output_dir.mkdir(parents=True, exist_ok=True)
        for name, content in rendered.items():
            target = args.output_dir / name
            temporary = target.with_suffix(target.suffix + ".tmp")
            temporary.write_text(content, encoding="utf-8")
            temporary.replace(target)
    print(json.dumps({
        **journeys["metadata"]["counts"],
        "placeCountsByLeg": dict(sorted(Counter(item["leg"] for item in places).items())),
        "placeSourceUrls": sum(len(item["sourceUrls"]) for item in places),
        "gapSourceUrls": sum(len(item["sourceUrls"]) for item in journeys["coverageGaps"]),
        "anonymizedCommentaryCells": journeys["metadata"]["anonymizedCommentaryCells"],
        "sourceUnchanged": before == after,
        "sourceSha256": before,
        "mode": "check" if args.check else "write",
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
